"""
Extracción de metadatos de archivos PowerPoint con python-pptx.

Solo soporta .pptx (Office Open XML). Los archivos .ppt legacy se validan
mediante LibreOffice durante la conversión.
"""
from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.exc import PackageNotFoundError


class PowerPointMetadataError(Exception):
    """Error al leer metadatos del PowerPoint."""


def extract_pptx_metadata(file_path: Path) -> dict:
    """
    Valida estructura PPTX y extrae títulos por diapositiva.
    Retorna: {"slide_count": int, "slides": [{"title": str, "index": int}, ...]}
    """
    try:
        presentation = PptxPresentation(str(file_path))
    except PackageNotFoundError as exc:
        raise PowerPointMetadataError(
            "El archivo está corrupto o no es un PowerPoint válido (.pptx)."
        ) from exc
    except Exception as exc:
        raise PowerPointMetadataError(
            f"No fue posible leer el archivo PowerPoint: {exc}"
        ) from exc

    slides_meta = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()
        slides_meta.append({"index": index, "title": title})

    if not slides_meta:
        raise PowerPointMetadataError("La presentación no contiene diapositivas.")

    return {
        "slide_count": len(slides_meta),
        "slides": slides_meta,
    }
